#!/usr/bin/env python3
"""
Generate the CFO presentation (Andrew Harper · Aldwych Bank) as a .pptx.

This mirrors, slide-for-slide, the customer-facing deck in
`cfo-presentation.html`. It is deliberately the "in-the-room" content only —
no internal prep material (no reconcile / to-validate / who-says-what).

Usage
-----
    pip install python-pptx        # one-time
    python3 generate_cfo_pptx.py   # writes cfo-presentation.pptx

Keep this in sync with cfo-presentation.html if the numbers change.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- palette (matches the HTML deck) -------------------------------------
NAVY      = RGBColor(0x0B, 0x1A, 0x2B)
NAVY_2    = RGBColor(0x0F, 0x29, 0x42)
GOLD      = RGBColor(0xD4, 0xAF, 0x37)
SKY       = RGBColor(0x29, 0xB5, 0xE8)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
INK_LIGHT = RGBColor(0xD7, 0xE2, 0xEE)
MUTED     = RGBColor(0x9F, 0xB3, 0xC8)

# 16:9 canvas
EMU_W = Inches(13.333)
EMU_H = Inches(7.5)

FONT = "Helvetica Neue"


def build():
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    blank = prs.slide_layouts[6]

    # ---- small helpers ---------------------------------------------------
    def slide():
        s = prs.slides.add_slide(blank)
        _bg(s, NAVY)
        _accent_bar(s)
        return s

    def _bg(s, color):
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = color

    def _accent_bar(s):
        bar = s.shapes.add_shape(
            1, Emu(0), Emu(0), Inches(0.09), EMU_H)  # 1 = rectangle
        bar.fill.solid()
        bar.fill.fore_color.rgb = GOLD
        bar.line.fill.background()

    def box(s, l, t, w, h):
        tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        return tb, tf

    def para(tf, text, size, color, bold=False, first=False,
             align=PP_ALIGN.LEFT, space_after=6, upper=False, spacing=0.0):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = text.upper() if upper else text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        return p

    def panel(s, l, t, w, h, fill=RGBColor(0x11, 0x24, 0x38),
              line=GOLD, line_w=0.75):
        sh = s.shapes.add_shape(5, Inches(l), Inches(t), Inches(w), Inches(h))  # 5 = rounded rect
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
        sh.shadow.inherit = False
        return sh

    def kicker(s, text):
        _, tf = box(s, 0.55, 0.5, 12, 0.5)
        para(tf, text, 12, GOLD, bold=True, first=True, upper=True, spacing=0.2)

    def title(s, text, top=0.95, size=34):
        _, tf = box(s, 0.55, top, 12.2, 1.4)
        para(tf, text, size, WHITE, bold=True, first=True)

    def footer(s):
        _, tf = box(s, 0.55, 7.05, 8, 0.35)
        para(tf, "Aldwych Bank · CFO presentation", 10, MUTED, first=True)

    # =====================================================================
    # SLIDE 1 · COVER
    # =====================================================================
    s = slide()
    _bg(s, NAVY)
    _, tf = box(s, 0.55, 0.5, 10, 0.5)
    para(tf, "Snowflake on Azure   ·   with Accenture", 13, INK_LIGHT, first=True)

    _, tf = box(s, 0.55, 2.05, 12.0, 2.6)
    para(tf, "Trusted, auditable regulatory reporting", 46, WHITE, bold=True, first=True, space_after=2)
    para(tf, "in days, not weeks.", 46, GOLD, bold=True, space_after=10)

    _, tf = box(s, 0.55, 4.55, 11.4, 1.6)
    para(tf,
         "Close the lineage gap that failed the audit — clear the FCA notice by Q4, "
         "unstick a £340M SME loan book, and free regulatory capital — on one governed "
         "platform across all seven systems, without leaving Azure.",
         18, INK_LIGHT, first=True)

    _, tf = box(s, 0.55, 6.35, 7, 1)
    para(tf, "Andrew Harper", 16, WHITE, bold=True, first=True, space_after=1)
    para(tf, "Chief Financial Officer · Aldwych Bank", 12, MUTED)

    _, tf = box(s, 8.3, 6.35, 4.45, 1)
    para(tf, "Grounded in conversations across your bank", 11, MUTED, first=True,
         align=PP_ALIGN.RIGHT, space_after=1)
    para(tf, "CRO · CDAIO · CTO · Data Architecture · Model Risk", 11, MUTED,
         align=PP_ALIGN.RIGHT)

    # =====================================================================
    # SLIDE 2 · FOUR QUESTIONS (2x2 quadrant)
    # =====================================================================
    s = slide()
    kicker(s, "The CFO's lens")
    title(s, "Four questions — answered")

    quad = [
        ("Can I trust the numbers?",
         "End-to-end column-level lineage from every regulatory figure back to source — "
         "the BCBS 239 gap the current estate can't close."),
        ("Does this reduce business risk?",
         "Clears the FCA notice, removes the £35–55M penalty tail, and puts model "
         "governance on one auditable foundation."),
        ("Can we operate more efficiently?",
         "The three-week manual reporting cycle collapses toward days, freeing the 55% "
         "of risk-engineering effort lost to legacy pipelines."),
        ("Will this drive profitable growth?",
         "Unsticks the £340M SME book and reinstates its model, frees capital to lend, "
         "and opens governed AI self-service inside two quarters."),
    ]
    cw, ch, gx, gy = 5.9, 2.35, 0.35, 0.3
    x0, y0 = 0.55, 2.15
    for idx, (q, a) in enumerate(quad):
        r, c = divmod(idx, 2)
        l = x0 + c * (cw + gx)
        t = y0 + r * (ch + gy)
        panel(s, l, t, cw, ch)
        _, tf = box(s, l + 0.3, t + 0.22, cw - 0.6, ch - 0.4)
        para(tf, q, 17, WHITE, bold=True, first=True, space_after=6)
        para(tf, a, 13.5, INK_LIGHT)
    footer(s)

    # =====================================================================
    # SLIDE 3 · FINANCIAL CASE (3 metrics + table)
    # =====================================================================
    s = slide()
    kicker(s, "Going high · the financial case")
    title(s, "Your own numbers")

    metrics = [
        ("£340M", "SME loan book stuck in manual review"),
        ("£35–55M", "FCA penalty exposure avoided"),
        ("£4.2M/yr", "manual credit-review labour"),
    ]
    mw, gx = 3.95, 0.28
    x0, t = 0.55, 2.05
    for idx, (n, l) in enumerate(metrics):
        lft = x0 + idx * (mw + gx)
        panel(s, lft, t, mw, 1.55, fill=RGBColor(0x1A, 0x22, 0x18), line=GOLD)
        _, tf = box(s, lft, t + 0.2, mw, 1.2)
        para(tf, n, 34, GOLD, bold=True, first=True, align=PP_ALIGN.CENTER, space_after=4)
        para(tf, l, 12, INK_LIGHT, align=PP_ALIGN.CENTER)

    rows = [
        ("Risk-engineering capacity lost to legacy pipelines", "55% of the team's week"),
        ("Regulatory reporting cycle", "3 weeks → days"),
        ("Model reinstatement P&L uplift", "~£10.3M  (CDO's estimate)"),
        ("Regulatory capital tied up under Basel IV", "freed by lineage → redeployed to lending"),
    ]
    _table(s, rows, left=0.55, top=3.95, width=12.2,
           col0=7.4, RGBColor=RGBColor, WHITE=WHITE, INK_LIGHT=INK_LIGHT,
           GOLD=GOLD, MUTED=MUTED, FONT=FONT, box=box, para=para, panel=panel)

    _, tf = box(s, 0.55, 6.55, 12.2, 0.6)
    para(tf,
         "The two rock-solid figures — £4.2M/yr labour and the £35–55M penalty tail — "
         "anchor the case; the £340M book is what the fix unsticks.",
         12, MUTED, first=True)

    # =====================================================================
    # SLIDE 4 · THE FIX PAYS FOR ITSELF
    # =====================================================================
    s = slide()
    kicker(s, "Cost of action vs inaction")
    title(s, "The fix pays for itself")

    _, tf = box(s, 0.55, 1.95, 12.0, 0.6)
    para(tf, "Judge it as total cost, not licence-versus-licence — the whole current "
             "estate against one governed platform.", 15, INK_LIGHT, first=True)

    rows = [
        ("Consolidate", "Seven systems, plus £4.2M/yr manual labour and maintenance drag, "
                        "collapse into one platform — retiring duplicate infrastructure and licences."),
        ("Self-funding", "The first step is largely a reallocation of budget already set aside — "
                         "payback in months, not years against £4.2M/yr recovered."),
        ("Capital freed", "Traceable lineage lets you retire conservative Basel IV assumptions — "
                          "capital redeployed to lending."),
        ("Cost of delay", "Every month the model stays off ≈ ~£0.86M foregone P&L + ~£0.35M labour. "
                          "Waiting isn't free."),
        ("Avoided rework", "One foundation satisfies BCBS 239, PRA SS1/23 and Basel IV — "
                           "not three separate remediation projects later."),
    ]
    _table(s, rows, left=0.55, top=2.75, width=12.2,
           col0=2.7, RGBColor=RGBColor, WHITE=WHITE, INK_LIGHT=INK_LIGHT,
           GOLD=GOLD, MUTED=MUTED, FONT=FONT, box=box, para=para, panel=panel,
           col0_bold=True)
    footer(s)

    # =====================================================================
    # SLIDE 5 · ONE CONNECTED ESTATE
    # =====================================================================
    s = slide()
    kicker(s, "Going wide · one platform, on Azure")
    title(s, "One connected estate — not a point fix")

    _, tf = box(s, 0.55, 2.1, 7.1, 4.6)
    bullets = [
        ("The gap: ", "the current tooling tracks assets, not the transformation-level "
                      "column lineage BCBS 239 demands — the same gap that failed the audit."),
        ("Snowflake on Azure ", "closes it now, preserving your Microsoft commitment and "
                                "Azure consumption. Coexistence, not rip-and-replace."),
        ("Seven systems → one governed platform, ", "with the ring-fenced estates kept as "
                                "separate tenants, unified only via governed data sharing — not merged."),
        ("Accenture co-delivers, ", "bridging the legacy estate and de-risking against "
                                "another mid-stream failure."),
    ]
    first = True
    for lead, rest in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(12)
        r0 = p.add_run(); r0.text = "•  " + lead
        r0.font.name = FONT; r0.font.size = Pt(15); r0.font.bold = True; r0.font.color.rgb = WHITE
        r1 = p.add_run(); r1.text = rest
        r1.font.name = FONT; r1.font.size = Pt(15); r1.font.color.rgb = INK_LIGHT

    panel(s, 8.05, 2.1, 4.7, 4.55, fill=RGBColor(0x11, 0x24, 0x38), line=GOLD)
    _, tf = box(s, 8.35, 2.35, 4.1, 4.1)
    para(tf, "ONE FOUNDATION, THREE OUTCOMES", 13, GOLD, bold=True, first=True, space_after=12)
    outcomes = [
        ("Regulatory reporting", "BCBS 239 + Basel IV capital relief"),
        ("Model reinstatement", "FCA / PRA SS1/23 governance"),
        ("Governed AI self-service", "plain-English queries for execs, analytics for the team"),
    ]
    for head, sub in outcomes:
        para(tf, "•  " + head, 15, WHITE, bold=True, space_after=1)
        para(tf, "     " + sub, 12.5, MUTED, space_after=10)
    para(tf, "Adjacent upside: AML / fraud analytics and an enterprise knowledge graph for AI trust.",
         11.5, MUTED, space_after=0)
    footer(s)

    # =====================================================================
    # SLIDE 6 · SEE IT LIVE + THE ASK
    # =====================================================================
    s = slide()
    kicker(s, "Going fast · proof, then the ask")
    title(s, "See it live — then a funded first step")

    panel(s, 0.55, 2.1, 6.2, 4.55, fill=RGBColor(0x11, 0x24, 0x38), line=GOLD)
    _, tf = box(s, 0.85, 2.35, 5.6, 4.1)
    para(tf, "LIVE DEMO — FOUR CLICKS", 13, GOLD, bold=True, first=True, space_after=12)
    steps = [
        "Start on a regulatory figure in Sigma, your chosen BI.",
        "Click through column-level lineage back to source — the exact thing the current estate can't do, shown live.",
        "Flip to the reinstated SME credit model with provenance and human-in-the-loop review.",
        "Ask a plain-English question a CFO would actually ask — under strict access control.",
    ]
    for idx, st in enumerate(steps, 1):
        para(tf, f"{idx}.  {st}", 14, INK_LIGHT, space_after=10)

    panel(s, 7.05, 2.1, 5.7, 3.0, fill=RGBColor(0x0C, 0x2A, 0x3A), line=SKY, line_w=1.0)
    _, tf = box(s, 7.35, 2.3, 5.1, 2.7)
    para(tf, "THE FUNDED FIRST STEP", 12, SKY, bold=True, first=True, space_after=8)
    para(tf, "Extend the existing proof-of-concept to cover retail and investment banking with "
             "column-level lineage, source→BI, on Snowflake-on-Azure — reinstating the SME credit model.",
         13.5, WHITE, space_after=8)
    para(tf, "Funded from budget already ring-fenced, delivered inside two quarters, ahead of Q4. "
             "Proof before commitment.", 13.5, INK_LIGHT)

    _, tf = box(s, 7.05, 5.35, 5.7, 1.3)
    para(tf, "Speed proof: a live HSBC reference on Snowflake + Sigma; governance that stands up "
             "fast; prior delivery at comparable global banks.", 11.5, MUTED, first=True)
    footer(s)

    out = "cfo-presentation.pptx"
    prs.save(out)
    print(f"Wrote {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


def _table(s, rows, left, top, width, col0,
           RGBColor, WHITE, INK_LIGHT, GOLD, MUTED, FONT, box, para, panel,
           col0_bold=False):
    """Simple 2-column body table drawn as text rows with divider lines."""
    from pptx.util import Inches, Pt
    row_h = 0.62
    for idx, (a, b) in enumerate(rows):
        t = top + idx * row_h
        _, tfa = box(s, left, t, col0, row_h)
        para(tfa, a, 13, WHITE if col0_bold else INK_LIGHT, bold=col0_bold, first=True)
        _, tfb = box(s, left + col0 + 0.2, t, width - col0 - 0.2, row_h)
        para(tfb, b, 13, INK_LIGHT, first=True)
        # divider
        ln = s.shapes.add_shape(1, Inches(left), Inches(t + row_h - 0.06),
                                Inches(width), Pt(0.75))
        ln.fill.solid(); ln.fill.fore_color.rgb = RGBColor(0x2A, 0x3B, 0x4C)
        ln.line.fill.background()


if __name__ == "__main__":
    build()
